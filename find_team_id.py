from pptx import Presentation
import glob, os, re

paths = glob.glob(r"C:\Users\pooja\Downloads\*.pptx") + glob.glob(r"C:\Users\pooja\agrobridge\*.pptx")
hits = []
for f in paths:
    try:
        prs = Presentation(f)
    except Exception:
        continue
    blob = []
    for slide in prs.slides:
        for sh in slide.shapes:
            if sh.has_text_frame:
                blob.append(sh.text_frame.text)
            if sh.has_table:
                for row in sh.table.rows:
                    for c in row.cells:
                        blob.append(c.text)
    text = "\n".join(blob)
    m = re.findall(r"(?:Team\s*ID|TeamID)\s*[:\-]?\s*([A-Za-z0-9\-\/]+)", text, re.I)
    if m:
        hits.append((os.path.basename(f), m))
for h in hits:
    print(h)
print("searched", len(paths), "files")
