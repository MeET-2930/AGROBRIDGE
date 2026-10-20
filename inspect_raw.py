from pptx import Presentation
from pptx.util import Emu
import sys, os, zipfile, io
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

def dump(f, label):
    print("=" * 80)
    print(label, os.path.basename(f))
    prs = Presentation(f)
    # also raw XML text extract via zip
    z = zipfile.ZipFile(f)
    for si in range(1, len(prs.slides) + 1):
        name = f"ppt/slides/slide{si}.xml"
        if name not in z.namelist():
            continue
        xml = z.read(name).decode("utf-8", errors="replace")
        # extract all <a:t> text
        import re
        texts = re.findall(r"<a:t[^>]*>(.*?)</a:t>", xml)
        texts = [t for t in texts if t.strip()]
        print(f"  raw slide {si} texts ({len(texts)}):")
        for t in texts:
            print(f"    - {t[:150]}")
        # count pics
        pics = xml.count("<p:pic>")
        print(f"    [pictures in xml: {pics}]")
    z.close()

    # walk shapes including groups
    def walk(shapes, depth=0, si=None):
        for sh in shapes:
            ind = "  " * depth
            try:
                pos = f"({Emu(sh.left).inches:.2f},{Emu(sh.top).inches:.2f},{Emu(sh.width).inches:.2f}x{Emu(sh.height).inches:.2f})"
            except Exception:
                pos = "?"
            print(f"{ind}  [{si}] {sh.shape_type} {sh.name!r} {pos}")
            if sh.shape_type == 6:  # group
                walk(sh.shapes, depth + 1, si)
            if sh.has_text_frame and sh.text_frame.text.strip():
                print(f"{ind}    TEXT: {sh.text_frame.text.strip()[:200]!r}")

    for si, slide in enumerate(prs.slides, 1):
        print(f"  -- shapes slide {si} --")
        walk(slide.shapes, 0, si)

dump(r"C:\Users\pooja\Downloads\Sih-26032.pptx", "TEAM OFFICIAL")
dump(r"C:\Users\pooja\Downloads\final_edit_2.pptx", "FINAL_EDIT_2")
dump(r"C:\Users\pooja\Downloads\SIH2026-IDEA-Presentation-Format.pptx", "BLANK OFFICIAL")
