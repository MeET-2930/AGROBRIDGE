from pptx import Presentation
from pptx.util import Emu
import sys
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

f = r"C:\Users\pooja\Downloads\Sih-26032.pptx"
prs = Presentation(f)
print("SLIDE SIZE:", Emu(prs.slide_width).inches, "x", Emu(prs.slide_height).inches)
print("MASTERS:", len(prs.slide_masters))
for mi, m in enumerate(prs.slide_masters):
    print(f" master {mi} theme-ish layouts:", [l.name for l in m.slide_layouts])
    # master shapes
    for sh in m.shapes:
        info = f"  master shape: {sh.shape_type} name={sh.name!r}"
        if sh.has_text_frame and sh.text_frame.text.strip():
            info += f" text={sh.text_frame.text.strip()[:80]!r}"
        print(info)

for si, slide in enumerate(prs.slides, 1):
    print("=" * 70)
    print(f"SLIDE {si}  layout={slide.slide_layout.name!r}")
    for idx, sh in enumerate(slide.shapes):
        line = f"  [{idx}] {sh.shape_type} name={sh.name!r} ph={sh.is_placeholder}"
        try:
            line += f" pos=({Emu(sh.left).inches:.2f},{Emu(sh.top).inches:.2f}) size=({Emu(sh.width).inches:.2f}x{Emu(sh.height).inches:.2f})"
        except Exception:
            pass
        if sh.is_placeholder:
            line += f" ph_type={sh.placeholder_format.type} idx={sh.placeholder_format.idx}"
        print(line)
        if sh.has_text_frame:
            for pi, p in enumerate(sh.text_frame.paragraphs):
                t = p.text
                if t.strip():
                    sizes = []
                    bolds = []
                    colors = []
                    fonts = []
                    for r in p.runs:
                        sizes.append(r.font.size.pt if r.font.size else None)
                        bolds.append(r.font.bold)
                        try:
                            colors.append(str(r.font.color.rgb) if r.font.color and r.font.color.rgb else None)
                        except Exception:
                            colors.append(None)
                        fonts.append(r.font.name)
                    print(f"      p{pi}: {t[:120]!r} sizes={sizes} bold={bolds} colors={colors} fonts={fonts}")
        if sh.has_table:
            tbl = sh.table
            print(f"      TABLE {len(tbl.rows)}x{len(tbl.columns)}")
            for ri, row in enumerate(tbl.rows):
                cells = [c.text.replace(chr(10), " ")[:40] for c in row.cells]
                print(f"        row{ri}: {cells}")
        if sh.shape_type == 13:  # picture
            print(f"      PICTURE image={getattr(sh.image, 'filename', '?')}")
