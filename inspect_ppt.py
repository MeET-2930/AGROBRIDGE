from pptx import Presentation
from pptx.util import Emu
import os, re

files = [
    r"C:\Users\pooja\Downloads\Sih-26032.pptx",
    r"C:\Users\pooja\Downloads\SIH2026-IDEA-Presentation-Format.pptx",
    r"C:\Users\pooja\Downloads\final_edit_2.pptx",
    r"C:\Users\pooja\Downloads\base ppt.pptx",
    r"C:\Users\pooja\Downloads\AgroBridge_SIH2026_BaseFormat.pptx",
    r"C:\Users\pooja\agrobridge\AgroBridge_SIH_2026_Round_2_Submission.pptx",
]
for f in files:
    if not os.path.exists(f):
        print("MISSING:", f)
        continue
    print("=" * 80)
    print("FILE:", os.path.basename(f), "size=", os.path.getsize(f))
    try:
        prs = Presentation(f)
    except Exception as e:
        print("  ERROR:", e)
        continue
    w, h = prs.slide_width, prs.slide_height
    print("  slides=", len(prs.slides), " dims=", Emu(w).inches, "x", Emu(h).inches, "in")
    print("  masters=", len(prs.slide_masters))
    for mi, m in enumerate(prs.slide_masters):
        print("    master[%d] layouts=%d:" % (mi, len(m.slide_layouts)), [l.name for l in m.slide_layouts])
    all_text = []
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for sh in slide.shapes:
            if sh.has_text_frame:
                t = sh.text_frame.text.strip()
                if t:
                    texts.append(t.replace("\n", " | "))
            if sh.has_table:
                for row in sh.table.rows:
                    for c in row.cells:
                        if c.text.strip():
                            texts.append("TBL:" + c.text.strip().replace("\n", " "))
        joined = " || ".join(texts)
        all_text.append(joined)
        print("  SLIDE %d: %s" % (i, joined[:700]))
    blob = " ".join(all_text)
    m = re.findall(r"[Tt]eam\s*ID\W{0,5}\s*([A-Za-z0-9\-]+)", blob)
    if m:
        print("  *** TEAM ID FOUND:", m)
